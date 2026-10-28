"""
Export the API contract for frontend handoff.

Outputs (all into contracts/):
  - openapi.json          (machine-readable, FastAPI-generated)
  - api_reference.md      (human-readable, enriched)
  - api_reference.xlsx    (spreadsheet — Endpoints + Schemas)
  - types.ts              (TS types; real json-schema-to-typescript if
                          Node available, else fallback)
  - index.md              (landing page — base URL, auth, examples)
  - README.md             (how to consume the bundle)
  - version_manifest.json (generated_at, commit_sha, counts)

Usage (PowerShell):
    python scripts/export_api_contract.py
"""

from __future__ import annotations

import hashlib
import json
import shutil
import subprocess
import sys
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

from social_api.main import create_app  # noqa: E402

CONTRACTS_DIR = ROOT / "contracts"


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _git_commit_sha() -> str:
    """Return current git HEAD SHA, or 'unknown' if unavailable.

    Best-effort: not a git repo, no git installed, or exec failure all
    return 'unknown' — export continues.
    """
    try:
        out = subprocess.run(
            ["git", "rev-parse", "HEAD"],
            cwd=ROOT,
            capture_output=True,
            text=True,
            timeout=5,
            check=False,
        )
        if out.returncode == 0:
            return out.stdout.strip()
    except (OSError, subprocess.SubprocessError) as exc:
        # Narrow: git not installed, exec failed, timeout. NOT catching
        # KeyboardInterrupt / SystemExit / MemoryError — those should
        # propagate.
        import logging

        logging.getLogger(__name__).debug("git_commit_sha_unavailable: %s", type(exc).__name__)
    return "unknown"


def _write_json(path: Path, data: dict[str, Any]) -> None:
    path.write_text(
        json.dumps(data, indent=2, ensure_ascii=False, sort_keys=True),
        encoding="utf-8",
    )


def _md_escape(s: str) -> str:
    return s.replace("|", "\\|").replace("\n", " ").strip()


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------
def _derive_auth(op: dict[str, Any]) -> str:
    """Return 'public' | 'authenticated' | role list (from x-required-roles)."""
    roles = op.get("x-required-roles")
    if roles:
        return "authenticated · roles: " + ", ".join(roles)
    security = op.get("security")
    if security:
        return "authenticated"
    return "public"


def _derive_rate_limit(op: dict[str, Any]) -> str:
    rl = op.get("x-rate-limit")
    if rl:
        return str(rl)
    return "—"


def _derive_errors(op: dict[str, Any]) -> str:
    codes = sorted(op.get("responses", {}).keys(), key=str)
    bad = [c for c in codes if str(c).startswith(("4", "5"))]
    return ", ".join(str(c) for c in bad) if bad else "—"


def _example_curl(method: str, path: str, op: dict[str, Any]) -> str:
    """Generate a cURL snippet from the endpoint metadata."""
    auth = _derive_auth(op)
    headers = ['-H "Content-Type: application/json"']
    if auth != "public":
        headers.append('-H "Authorization: Bearer $TOKEN"')
    rl = _derive_rate_limit(op)
    if rl and rl != "—":
        headers.append(f'-H "# rate-limit: {rl}"')

    body = ""
    rb = op.get("requestBody", {}).get("content", {}).get("application/json")
    if rb:
        example: dict[str, Any] = {}
        schema = rb.get("schema", {})
        ref = schema.get("$ref", "")
        if ref:
            example["<see schema>"] = ref.rsplit("/", 1)[-1]
        if example:
            body = " -d '" + json.dumps(example) + "'"

    return (
        f"curl -X {method.upper()} {' '.join(headers)}{body} \\\n     http://localhost:8000{path}"
    )


def _write_markdown(path: Path, spec: dict[str, Any]) -> None:
    lines: list[str] = []
    lines.append(f"# {spec['info'].get('title', 'API')} — API Reference\n")
    lines.append(f"**Version:** `{spec['info'].get('version', '?')}`\n")
    lines.append(f"**Generated from commit:** `{_git_commit_sha()[:12]}`\n")
    lines.append("---\n")

    # Endpoints grouped by tag.
    # NOTE: loop var is path_url (not path) so it does not shadow the
    # `path: Path` parameter used at the end for .write_text().
    by_tag: dict[str, list[tuple[str, str, dict[str, Any]]]] = {}
    for path_url, methods in spec.get("paths", {}).items():
        for method, op in methods.items():
            if method.lower() not in {"get", "post", "put", "patch", "delete"}:
                continue
            tags = op.get("tags") or ["untagged"]
            primary = tags[0]
            by_tag.setdefault(primary, []).append((method, path_url, op))

    for tag in sorted(by_tag):
        lines.append(f"\n## {tag}\n")
        for method, path_url, op in sorted(by_tag[tag], key=lambda t: t[1]):
            summary = op.get("summary", "")
            lines.append(f"### `{method.upper()}` {path_url}\n")
            if summary:
                lines.append(f"**Summary:** {_md_escape(summary)}\n")
            lines.append(f"**Auth:** {_derive_auth(op)}\n")
            lines.append(f"**Rate limit:** {_derive_rate_limit(op)}\n")
            lines.append(f"**Errors:** {_derive_errors(op)}\n")

            params = op.get("parameters", [])
            if params:
                lines.append("\n**Parameters:**\n")
                lines.append("| Name | In | Required | Schema |")
                lines.append("|---|---|---|---|")
                for p in params:
                    schema_str = json.dumps(p.get("schema", {}), ensure_ascii=False)
                    lines.append(
                        f"| {p.get('name')} | {p.get('in')} | "
                        f"{p.get('required', False)} | "
                        f"`{_md_escape(schema_str)}` |"
                    )
                lines.append("")

            rb = op.get("requestBody")
            if rb:
                lines.append("\n**Request body:**\n")
                for ctype, media in rb.get("content", {}).items():
                    ref = media.get("schema", {}).get("$ref", "")
                    name = ref.rsplit("/", 1)[-1] if ref else "inline"
                    lines.append(f"- `{ctype}` -> `{name}`")
                lines.append("")

            if op.get("responses"):
                lines.append("\n**Responses:**\n")
                for code, resp in sorted(op["responses"].items(), key=lambda kv: str(kv[0])):
                    desc = resp.get("description", "")
                    lines.append(f"- `{code}` — {_md_escape(desc)}")
                lines.append("")

            lines.append("\n**Example:**\n")
            lines.append("```bash")
            lines.append(_example_curl(method, path_url, op))
            lines.append("```\n")
            lines.append("---\n")

    path.write_text("\n".join(lines), encoding="utf-8")


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def _write_excel(path: Path, spec: dict[str, Any]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        print("  ! openpyxl not installed - skipping XLSX")
        return

    wb = Workbook()

    ws = wb.active
    if ws is None:
        raise RuntimeError(
            "openpyxl Workbook has no active worksheet — openpyxl install "
            "is broken or a subclass overrode .active to return None."
        )

    ws.title = "Endpoints"

    headers = [
        "Method",
        "Path",
        "Summary",
        "Tags",
        "Auth",
        "Rate Limit",
        "Errors",
        "Request Schema",
        "Response Codes",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for path_url, methods in sorted(spec.get("paths", {}).items()):
        for method, op in sorted(methods.items()):
            if method.lower() not in {"get", "post", "put", "patch", "delete"}:
                continue
            request_ref = ""
            rb = op.get("requestBody")
            if rb:
                for media in rb.get("content", {}).values():
                    ref = media.get("schema", {}).get("$ref", "")
                    if ref:
                        request_ref = ref.rsplit("/", 1)[-1]
                        break
            response_codes = ", ".join(
                str(c) for c in sorted(op.get("responses", {}).keys(), key=str)
            )
            ws.append(
                [
                    method.upper(),
                    path_url,
                    op.get("summary", ""),
                    ", ".join(op.get("tags", [])),
                    _derive_auth(op),
                    _derive_rate_limit(op),
                    _derive_errors(op),
                    request_ref,
                    response_codes,
                ]
            )

    # Sheet 2: Schemas
    ws2 = wb.create_sheet("Schemas")
    ws2.append(["Name", "Type", "Required", "Properties"])
    for cell in ws2[1]:
        cell.font = Font(bold=True)

    for name, schema in sorted(spec.get("components", {}).get("schemas", {}).items()):
        required = ", ".join(schema.get("required", []))
        props = ", ".join(schema.get("properties", {}).keys())
        ws2.append([name, schema.get("type", ""), required, props])

    # Column widths (rough)
    for ws_i, widths in [
        (ws, [10, 40, 40, 20, 30, 15, 20, 30, 20]),
        (ws2, [30, 10, 30, 60]),
    ]:
        for col_idx, w in enumerate(widths, start=1):
            ws_i.column_dimensions[chr(64 + col_idx)].width = w

    wb.save(path)


# ---------------------------------------------------------------------------
# TypeScript
# ---------------------------------------------------------------------------
def _write_types_ts(spec: dict[str, Any], path: Path) -> None:
    """
    Prefer npx json-schema-to-typescript when Node is available.
    Otherwise use a fallback generator.
    """
    node = shutil.which("node")
    npx = shutil.which("npx")
    if node and npx:
        try:
            result = subprocess.run(
                [
                    "npx",
                    "--yes",
                    "json-schema-to-typescript",
                    "-i",
                    str(CONTRACTS_DIR / "openapi.json"),
                    "-o",
                    str(path),
                    "--bannerComment",
                    "// AUTO-GENERATED by scripts/export_api_contract.py",
                ],
                timeout=120,
                capture_output=True,
                check=False,
            )
            if result.returncode == 0 and path.exists():
                print("      (via npx json-schema-to-typescript)")
                return
            print("      ! npx failed; falling back to Python generator")
        except Exception as exc:
            # Broad catch is intentional: npx is best-effort. Any problem
            # (broken install, exec perm, network) should fall through to
            # the Python generator rather than fail the whole export.
            print(f"      ! npx error ({exc}); falling back")

    # Fallback generator
    components = spec.get("components", {}).get("schemas", {})
    lines = [
        "// AUTO-GENERATED by scripts/export_api_contract.py",
        "// Source: contracts/openapi.json",
        "// NOTE: fallback generator (run with Node to get full fidelity)",
        "",
    ]

    def _ts_type(schema: dict[str, Any]) -> str:
        if "$ref" in schema:
            return str(schema["$ref"]).rsplit("/", 1)[-1]
        t = schema.get("type")
        if t == "string":
            return "string"
        if t in ("integer", "number"):
            return "number"
        if t == "boolean":
            return "boolean"
        if t == "null":
            return "null"
        if t == "array":
            return f"{_ts_type(schema.get('items', {}))}[]"
        if t == "object":
            add = schema.get("additionalProperties")
            if isinstance(add, dict):
                return f"Record<string, {_ts_type(add)}>"
            return "Record<string, unknown>"
        if "anyOf" in schema:
            return " | ".join(_ts_type(s) for s in schema["anyOf"])
        if "oneOf" in schema:
            return " | ".join(_ts_type(s) for s in schema["oneOf"])
        return "unknown"

    for name, schema in sorted(components.items()):
        props = schema.get("properties", {})
        required = set(schema.get("required", []))
        if not props:
            lines.append(f"export type {name} = {_ts_type(schema)};\n")
            continue
        lines.append(f"export interface {name} {{")
        for pname, pschema in props.items():
            optional = "" if pname in required else "?"
            lines.append(f"  {pname}{optional}: {_ts_type(pschema)};")
        lines.append("}\n")

    path.write_text("\n".join(lines), encoding="utf-8")


# ---------------------------------------------------------------------------
# index.md
# ---------------------------------------------------------------------------
def _write_index_md(path: Path, spec: dict[str, Any], manifest: dict[str, Any]) -> None:
    body = (
        f"# API Contract - {spec['info'].get('title', 'API')}\n\n"
        f"**Version:** `{spec['info'].get('version', '?')}`  \n"
        f"**Generated from commit:** `{manifest['commit_sha'][:12]}` on "
        f"{manifest['generated_at']}\n\n"
        "---\n\n"
        "## Quick start\n\n"
        "| Artifact | Use it for |\n"
        "|---|---|\n"
        "| `openapi.json` | Code generation (openapi-typescript, openapi-generator) |\n"
        "| `api_reference.md` | Human-readable per-endpoint documentation |\n"
        "| `api_reference.xlsx` | Spreadsheet review with backend/PM |\n"
        "| `types.ts` | Direct import of TypeScript types |\n"
        "| `README.md` | How to consume this bundle |\n\n"
        "## Base URL\n\n"
        "`http://localhost:8000/api/v1` (dev)\n\n"
        "## Authentication\n\n"
        "All authenticated endpoints require:\n\n"
        "```\n"
        "Authorization: Bearer <access_token>\n"
        "```\n\n"
        "Get a token via `POST /api/v1/auth/login` (returns "
        "`{tokens: {access_token, refresh_token}, user: {...}}`).\n\n"
        "Access token lifetime: 15 minutes. Refresh via "
        "`POST /api/v1/auth/refresh`.\n\n"
        "## Error envelope\n\n"
        "All errors (4xx/5xx) share one JSON shape:\n\n"
        "```json\n"
        "{\n"
        '  "error": "POST_NOT_FOUND",\n'
        '  "status": 404,\n'
        '  "detail": "Post not found",\n'
        '  "path": "/api/v1/posts/123",\n'
        '  "request_id": "8b6e3cf1-..."\n'
        "}\n"
        "```\n\n"
        "The `request_id` matches the `X-Request-ID` response header - "
        "include it when reporting bugs.\n\n"
        "## Response headers\n\n"
        "Every response includes:\n\n"
        "- `X-Request-ID` - unique per request\n"
        "- `X-Process-Time-Ms` - server-side processing duration\n\n"
        "## Cursor pagination\n\n"
        "List endpoints use cursor pagination, not OFFSET:\n\n"
        "```\n"
        "GET /api/v1/posts?limit=20&cursor=<opaque>\n"
        "```\n\n"
        "`response.next_cursor` is a stable opaque token - pass it back "
        "unchanged to fetch the next page. `next_cursor: null` means end "
        "of data.\n\n"
        "## Rate limits\n\n"
        "Endpoints with `x-rate-limit` metadata enforce per-IP or "
        "per-user limits. 429 responses include a `Retry-After` header "
        "(seconds).\n\n"
        "## Idempotency\n\n"
        "Write endpoints that accept `Idempotency-Key: <uuid>` will:\n\n"
        "- **First call**: execute normally\n"
        "- **Replay (same key, same body)**: return the cached response "
        "with `Idempotent-Replay: true` header\n"
        "- **Conflict (same key, different body)**: 409 "
        "`IDEMPOTENCY_KEY_CONFLICT`\n\n"
        "Replays are safe to retry - the underlying resource is not "
        "created twice.\n"
    )
    path.write_text(body, encoding="utf-8")


# ---------------------------------------------------------------------------
# README.md
# ---------------------------------------------------------------------------
def _write_readme(path: Path, manifest: dict[str, Any]) -> None:
    body = (
        "# API Contract Bundle\n\n"
        f"Auto-generated from commit `{manifest['commit_sha'][:12]}` at "
        f"{manifest['generated_at']}.\n\n"
        "**Do not edit these files by hand.** Regenerate with:\n\n"
        "```powershell\n"
        "python scripts/export_api_contract.py\n"
        "```\n\n"
        "## What's here\n\n"
        "| File | Purpose |\n"
        "|---|---|\n"
        "| `openapi.json` | OpenAPI 3.1 spec - machine-readable source of truth |\n"
        "| `api_reference.md` | Human-readable endpoint docs |\n"
        "| `api_reference.xlsx` | Spreadsheet view for reviews |\n"
        "| `types.ts` | TypeScript types generated from the OpenAPI schemas |\n"
        "| `index.md` | Landing page - base URL, auth, conventions |\n"
        "| `version_manifest.json` | Metadata: generated_at, commit SHA, counts |\n\n"
        "## How to consume (frontend project)\n\n"
        "### TypeScript (recommended)\n\n"
        "Import types directly:\n\n"
        "```ts\n"
        "import type { Post, User, Comment } from './contracts/types';\n"
        "```\n\n"
        "Or generate a client with `openapi-typescript`:\n\n"
        "```bash\n"
        "npx openapi-typescript ./contracts/openapi.json -o ./src/api/schema.ts\n"
        "```\n\n"
        "### Fetch client (framework-agnostic)\n\n"
        "```ts\n"
        "const res = await fetch(`${API_BASE}/api/v1/posts`, {\n"
        "  headers: { Authorization: `Bearer ${token}` },\n"
        "});\n"
        "if (!res.ok) {\n"
        "  const err = await res.json();\n"
        "  throw new ApiError(err.error, err.status, err.detail, err.request_id);\n"
        "}\n"
        "```\n\n"
        "## When this changes\n\n"
        "- **Breaking change** -> new `/api/v2/` endpoint, v1 kept alive 12 months\n"
        "- **Additive change** (new optional field, new endpoint) -> no "
        "version bump, regenerate bundle\n"
        "- **Deprecation** -> `Deprecation` + `Sunset` headers on affected "
        "endpoints\n\n"
        "## Drift prevention\n\n"
        "CI runs `scripts/check_api_contract_drift.py` on every push. If "
        "`contracts/openapi.json` is stale relative to `src/`, CI fails. "
        "Regenerate locally before committing API changes.\n"
    )
    path.write_text(body, encoding="utf-8")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main() -> None:
    print("[1/7] Generating openapi.json ...")
    app = create_app()
    spec = app.openapi()
    CONTRACTS_DIR.mkdir(exist_ok=True)
    openapi_path = CONTRACTS_DIR / "openapi.json"
    _write_json(openapi_path, spec)
    openapi_bytes = openapi_path.read_bytes()
    print(f"      -> {openapi_path}")

    print("[2/7] Generating api_reference.md ...")
    md_path = CONTRACTS_DIR / "api_reference.md"
    _write_markdown(md_path, spec)
    print(f"      -> {md_path}")

    print("[3/7] Generating api_reference.xlsx ...")
    xlsx_path = CONTRACTS_DIR / "api_reference.xlsx"
    _write_excel(xlsx_path, spec)
    print(f"      -> {xlsx_path}")

    print("[4/7] Generating types.ts ...")
    ts_path = CONTRACTS_DIR / "types.ts"
    _write_types_ts(spec, ts_path)
    print(f"      -> {ts_path}")

    print("[5/7] Generating index.md ...")
    endpoints_count = sum(
        len([m for m in methods if m.lower() in {"get", "post", "put", "patch", "delete"}])
        for methods in spec.get("paths", {}).values()
    )
    schemas_count = len(spec.get("components", {}).get("schemas", {}))
    manifest: dict[str, Any] = {
        "generated_at": datetime.now(UTC).isoformat(),
        "commit_sha": _git_commit_sha(),
        "openapi_sha256": _sha256(openapi_bytes),
        "endpoints_count": endpoints_count,
        "schemas_count": schemas_count,
    }
    _write_index_md(CONTRACTS_DIR / "index.md", spec, manifest)
    print(f"      -> {CONTRACTS_DIR / 'index.md'}")

    print("[6/7] Generating README.md ...")
    _write_readme(CONTRACTS_DIR / "README.md", manifest)
    print(f"      -> {CONTRACTS_DIR / 'README.md'}")

    print("[7/7] Writing version_manifest.json ...")
    _write_json(CONTRACTS_DIR / "version_manifest.json", manifest)
    print(f"      -> {CONTRACTS_DIR / 'version_manifest.json'}")

    print("\nDone. Contracts written to ./contracts/")


if __name__ == "__main__":
    main()
